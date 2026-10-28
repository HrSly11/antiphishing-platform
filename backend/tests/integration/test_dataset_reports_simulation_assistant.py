"""Integration tests for Dataset ETL idempotency, PDF/DOCX/XLSX Reports, Simulation, & Assistant."""
from __future__ import annotations

import io

import docx
import httpx
import openpyxl


async def test_dataset_etl_idempotency_and_sha256_lineage(
    client: httpx.AsyncClient,
    admin_headers: dict[str, str],
) -> None:
    # First run inserts deduplicated phishing URL samples from offline OpenPhish fixture
    run1 = await client.post(
        "/api/v1/dataset/ingest",
        headers=admin_headers,
        json={"source": "openphish_sample", "prefer_live": False, "sync_blacklist": True},
    )
    assert run1.status_code == 200
    v1 = run1.json()
    assert v1["status"] == "COMPLETED"
    assert v1["total_rows"] > 0
    assert v1["new_rows"] > 0
    assert len(v1["sha256"]) == 64

    # Second run over the exact same feed is idempotent: 0 new rows, >0 duplicates, identical SHA-256
    run2 = await client.post(
        "/api/v1/dataset/ingest",
        headers=admin_headers,
        json={"source": "openphish_sample", "prefer_live": False, "sync_blacklist": True},
    )
    assert run2.status_code == 200
    v2 = run2.json()
    assert v2["new_rows"] == 0
    assert v2["duplicate_rows"] == v1["total_rows"]
    assert v2["sha256"] == v1["sha256"]

    info_resp = await client.get("/api/v1/dataset/info", headers=admin_headers)
    assert info_resp.status_code == 200
    assert info_resp.json()["total_samples"] >= v1["new_rows"]

    samples_resp = await client.get("/api/v1/dataset/samples?page_size=10", headers=admin_headers)
    assert samples_resp.status_code == 200
    assert samples_resp.json()["total"] >= v1["new_rows"]


async def test_binary_reports_pdf_docx_xlsx_and_schedule_crud(
    client: httpx.AsyncClient,
    socl2_headers: dict[str, str],
) -> None:
    # 1. GET /reports/export/pdf -> valid %PDF- binary
    pdf_resp = await client.get(
        "/api/v1/reports/export/pdf?template=executive&window_days=30",
        headers=socl2_headers,
    )
    assert pdf_resp.status_code == 200
    assert "application/pdf" in pdf_resp.headers.get("content-type", "")
    assert pdf_resp.content.startswith(b"%PDF-")
    assert len(pdf_resp.content) > 1000

    # 2. GET /reports/export/docx -> parseable Word document
    docx_resp = await client.get(
        "/api/v1/reports/export/docx?template=management&window_days=30",
        headers=socl2_headers,
    )
    assert docx_resp.status_code == 200
    parsed_doc = docx.Document(io.BytesIO(docx_resp.content))
    full_doc_text = "\n".join(p.text for p in parsed_doc.paragraphs)
    assert len(full_doc_text) > 50
    assert len(parsed_doc.tables) >= 1

    # 3. GET /reports/export/xlsx -> parseable Excel workbook
    xlsx_resp = await client.get(
        "/api/v1/reports/export/xlsx?template=compliance&window_days=30",
        headers=socl2_headers,
    )
    assert xlsx_resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_resp.content))
    assert len(wb.sheetnames) >= 2
    active_ws = wb[wb.sheetnames[0]]
    assert active_ws.max_row >= 2

    # 4. POST /reports/generate JSON endpoint
    gen_resp = await client.post(
        "/api/v1/reports/generate",
        headers=socl2_headers,
        json={"report_format": "pdf", "template": "operational", "window_days": 14},
    )
    assert gen_resp.status_code == 200
    assert gen_resp.content.startswith(b"%PDF-")

    # 5. Report schedule CRUD (create -> list -> run-due -> delete)
    sched_create = await client.post(
        "/api/v1/reports/schedules",
        headers=socl2_headers,
        json={
            "name": "Daily SOC Morning Executive Digest",
            "report_format": "pdf",
            "frequency": "DAILY",
            "recipients": ["ciso@bancocentral.com", "socl2@bancocentral.com"],
            "template": "executive",
        },
    )
    assert sched_create.status_code == 201
    sched_id = sched_create.json()["id"]

    sched_list = await client.get("/api/v1/reports/schedules", headers=socl2_headers)
    assert sched_list.status_code == 200
    assert any(s["id"] == sched_id for s in sched_list.json())

    run_due = await client.post(
        "/api/v1/reports/schedules/run-due?force_all=true",
        headers=socl2_headers,
    )
    assert run_due.status_code == 200

    sched_del = await client.delete(
        f"/api/v1/reports/schedules/{sched_id}",
        headers=socl2_headers,
    )
    assert sched_del.status_code == 200
    assert sched_del.json()["status"] == "deleted"


async def test_phishing_simulation_campaign_events_and_rates(
    client: httpx.AsyncClient,
    socl2_headers: dict[str, str],
) -> None:
    camp_resp = await client.post(
        "/api/v1/simulation/campaigns",
        headers=socl2_headers,
        json={
            "name": "Treasury Wire Verification Awareness Drill",
            "channel": "EMAIL",
            "template_subject": "Accion requerida: Actualice su firma SWIFT",
            "lure_type": "CREDENTIAL_HARVEST",
            "target_department": "Treasury",
            "total_targets": 10,
        },
    )
    assert camp_resp.status_code == 201
    camp_id = camp_resp.json()["id"]

    events = [
        ("emp1@bancocentral.com", "SENT"),
        ("emp2@bancocentral.com", "SENT"),
        ("emp3@bancocentral.com", "SENT"),
        ("emp4@bancocentral.com", "SENT"),
        ("emp1@bancocentral.com", "OPENED"),
        ("emp2@bancocentral.com", "OPENED"),
        ("emp1@bancocentral.com", "CLICKED"),
        ("emp2@bancocentral.com", "REPORTED"),
        ("emp3@bancocentral.com", "REPORTED"),
    ]
    latest_campaign = None
    for emp_email, ev_type in events:
        ev_resp = await client.post(
            f"/api/v1/simulation/campaigns/{camp_id}/events",
            headers=socl2_headers,
            json={"employee_email": emp_email, "event_type": ev_type, "department": "Treasury"},
        )
        assert ev_resp.status_code == 201
        latest_campaign = ev_resp.json()["campaign"]

    assert latest_campaign is not None
    assert latest_campaign["sent_count"] == 4
    assert latest_campaign["opened_count"] == 2
    assert latest_campaign["clicked_count"] == 1
    assert latest_campaign["reported_count"] == 2
    assert latest_campaign["click_rate"] > 0.0
    assert latest_campaign["report_rate"] > 0.0

    metrics_resp = await client.get("/api/v1/simulation/metrics", headers=socl2_headers)
    assert metrics_resp.status_code == 200
    assert metrics_resp.json()["total_campaigns"] >= 1


async def test_soc_assistant_rag_citations_url_agent_and_langflow_sandboxes(
    client: httpx.AsyncClient,
    socl1_headers: dict[str, str],
) -> None:
    # 1. Grounded RAG assistant chat returning docs/ citations + live telemetry
    chat_resp = await client.post(
        "/api/v1/assistant/chat",
        headers=socl1_headers,
        json={
            "question": "What are the incident response containment steps and PCI-DSS audit controls?",
            "top_k": 4,
        },
    )
    assert chat_resp.status_code == 200
    chat_data = chat_resp.json()
    assert len(chat_data["answer"]) > 40
    assert len(chat_data["citations"]) >= 1
    assert all(c.startswith("docs/") for c in chat_data["citations"])
    assert chat_data["live_telemetry"]["connected"] is True

    # 2. Multi-tool LangChain URL investigation agent
    url_resp = await client.post(
        "/api/v1/assistant/investigate-url",
        headers=socl1_headers,
        json={
            "url": "http://bancocentral-seguridad.com/login/token",
            "enable_network_tools": False,
        },
    )
    assert url_resp.status_code == 200
    url_data = url_resp.json()
    assert url_data["verdict"] == "MALICIOUS"
    assert url_data["action"] == "QUARANTINE"
    assert url_data["risk_score"] >= 85
    assert "blacklist_tool" in url_data["tool_observations"]
    assert "pg_trgm_tool" in url_data["tool_observations"]

    # 3. Exported LangFlow v1 visual sandbox flows
    flows_resp = await client.get("/api/v1/assistant/flows", headers=socl1_headers)
    assert flows_resp.status_code == 200
    flows_data = flows_resp.json()
    assert flows_data["count"] == 2
    flow_names = {f["name"] for f in flows_data["flows"]}
    assert flow_names == {"email_classifier_sandbox", "url_agent_sandbox"}
