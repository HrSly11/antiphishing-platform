"""Word (.docx) corporate security report renderer using `python-docx`."""
from __future__ import annotations

import io
from datetime import datetime, timezone
from typing import Any, Dict, List

from docx import Document
from docx.shared import Pt, RGBColor

from app.adapters.reporting.pdf_renderer import TEMPLATE_TITLES


def render_docx_report(report_data: Dict[str, Any]) -> bytes:
    """Build a structured `.docx` report for executive, operational, or compliance audiences."""
    template = str(report_data.get("template") or "operational").lower()
    title = str(
        report_data.get("title")
        or TEMPLATE_TITLES.get(template, TEMPLATE_TITLES["operational"])
    )
    date_from = str(report_data.get("date_from") or "Last 30 Days")
    date_to = str(report_data.get("date_to") or datetime.now(timezone.utc).strftime("%Y-%m-%d"))
    kpis = report_data.get("kpis") or {}
    sla = report_data.get("sla_metrics") or {}
    verdicts: List[Dict[str, Any]] = list(report_data.get("verdicts") or [])
    incidents: List[Dict[str, Any]] = list(report_data.get("incidents") or [])
    top_iocs: List[Dict[str, Any]] = list(report_data.get("top_iocs") or [])

    doc = Document()
    doc.core_properties.author = "Banco Central SOC Antiphishing Platform"
    doc.core_properties.title = title

    heading = doc.add_heading(title, level=0)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(15, 23, 42)

    meta_p = doc.add_paragraph()
    meta_run = meta_p.add_run(
        f"Institution: Banco Central Corporate SOC | Window: {date_from} to {date_to} | "
        f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')} | "
        f"Template: {template.upper()} | Classification: CONFIDENTIAL"
    )
    meta_run.font.size = Pt(9)

    # 1. KPI Summary
    doc.add_heading("1. Key Performance & SLA Metrics", level=1)
    by_verdict = kpis.get("by_verdict") or {}
    kpi_table = doc.add_table(rows=1, cols=2)
    kpi_table.style = "Table Grid"
    hdr = kpi_table.rows[0].cells
    hdr[0].text = "Security / Operational Metric"
    hdr[1].text = "Observed Value"

    metrics_list = [
        ("Total Analyzed Messages", str(kpis.get("total_analyzed", len(verdicts)))),
        (
            "Malicious / Quarantined Verdicts",
            str(by_verdict.get("MALICIOUS", by_verdict.get("PHISHING", 0))),
        ),
        ("Suspicious Verdicts (Tagged)", str(by_verdict.get("SUSPICIOUS", 0))),
        ("Clean Verdicts (Delivered)", str(by_verdict.get("CLEAN", 0))),
        ("Average / P95 Inspection Latency (ms)", f"{kpis.get('latency_avg_ms', 0)} / {kpis.get('latency_p95_ms', 0)}"),
        ("False Positive / False Negative Rate", f"{kpis.get('false_positive_rate', 0.0)} / {kpis.get('false_negative_rate', 0.0)}"),
        ("Open / SLA Overdue Incidents", f"{sla.get('open_count', 0)} / {sla.get('overdue_count', 0)}"),
        (
            "SLA Compliance Rate / MTTR (minutes)",
            f"{round(float(sla.get('sla_compliance_rate', 1.0)) * 100, 1)}% / {sla.get('mttr_minutes') or 'N/A'}",
        ),
    ]
    for label, val in metrics_list:
        row_cells = kpi_table.add_row().cells
        row_cells[0].text = label
        row_cells[1].text = val

    # 2. Detection Verdicts
    doc.add_heading("2. Detection Verdicts Log", level=1)
    v_table = doc.add_table(rows=1, cols=6)
    v_table.style = "Table Grid"
    for idx, col_name in enumerate(["UTC Timestamp", "Sender", "Subject", "Channel", "Score", "Verdict"]):
        v_table.rows[0].cells[idx].text = col_name

    for v in verdicts[:30]:
        cells = v_table.add_row().cells
        cells[0].text = str(v.get("created_at") or v.get("timestamp") or "")[:19]
        cells[1].text = str(v.get("sender_email") or v.get("sender") or "")[:40]
        cells[2].text = str(v.get("subject") or "")[:50]
        cells[3].text = str(v.get("channel") or "EMAIL")
        cells[4].text = str(v.get("risk_score", 0))
        cells[5].text = str(v.get("verdict") or "")

    # 3. Incidents
    doc.add_heading("3. SOC Incident Register", level=1)
    i_table = doc.add_table(rows=1, cols=5)
    i_table.style = "Table Grid"
    for idx, col_name in enumerate(["Incident Code", "Severity", "Status", "Resolution", "SLA Overdue"]):
        i_table.rows[0].cells[idx].text = col_name

    for inc in incidents[:25]:
        cells = i_table.add_row().cells
        cells[0].text = str(inc.get("incident_code") or "")
        cells[1].text = str(inc.get("severity") or "")
        cells[2].text = str(inc.get("status") or "")
        cells[3].text = str(inc.get("resolution") or "PENDING")
        cells[4].text = "YES" if inc.get("sla_overdue") else "NO"

    # 4. Top IOCs & Compliance Mapping
    doc.add_heading("4. Active Threat Indicators & Regulatory Mapping", level=1)
    if top_iocs:
        ioc_table = doc.add_table(rows=1, cols=4)
        ioc_table.style = "Table Grid"
        for idx, col_name in enumerate(["Type", "Indicator", "Confidence", "Source"]):
            ioc_table.rows[0].cells[idx].text = col_name
        for ioc in top_iocs[:20]:
            cells = ioc_table.add_row().cells
            cells[0].text = str(ioc.get("ioc_type") or "")
            cells[1].text = str(ioc.get("ioc_value") or "")[:80]
            cells[2].text = str(ioc.get("confidence_score") or 100)
            cells[3].text = str(ioc.get("source") or "")

    doc.add_paragraph(
        "Regulatory Attestation: Evidence aligned with PCI-DSS v4.0 (Req 5.4.1, 10.2, 10.3, 12.10) "
        "and ISO/IEC 27001:2022 Annex A (A.5.7, A.5.24-A.5.26, A.8.7, A.8.15)."
    )

    buffer = io.BytesIO()
    doc.save(buffer)
    data = buffer.getvalue()
    buffer.close()
    return data
